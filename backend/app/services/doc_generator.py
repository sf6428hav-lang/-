from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def generate_docx(content: str, title: str, output_path: Path) -> Path:
    """生成排版后的 .docx 剧本文件。
    
    排版规则：
    - 全文统一使用黑体（SimHei / 微软雅黑），常规字形
    - 不添加斜体、不加粗等花哨样式
    - 剧名作为文档标题居中显示
    - 正文按原始换行输出，段前段后留适当间距方便阅读
    - 场次行、字幕行、动作行、台词行均保持统一常规样式
    """
    doc = Document()

    # 页面基础边距
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(3.0)

    # 默认字体：黑体族，中英文统一
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    pf = normal.paragraph_format
    pf.line_spacing = 1.5
    pf.space_after = Pt(0)

    # 文档标题（居中，稍大但保持常规不加粗）
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run(title)
    title_run.font.size = Pt(16)
    title_run.bold = False
    title_run.italic = False

    lines = content.split("\n")

    for idx, line in enumerate(lines):
        raw = line.rstrip()
        if not raw.strip():
            # 空行：给一个空白段落，保持可读间距
            doc.add_paragraph("")
            continue

        p = doc.add_paragraph()
        run = p.add_run(raw)
        run.bold = False
        run.italic = False
        run.font.size = Pt(11)

        # 场次行（形如 “1-1 日/内 客厅”）上方加一点间距，作为视觉分节
        if len(raw) > 2 and raw[0].isdigit() and "-" in raw[:5]:
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)

    doc.save(str(output_path))
    return output_path


def generate_txt(content: str, output_path: Path) -> Path:
    """保存纯文本（按原始换行输出，保证排版一致）。"""
    output_path.write_text(content, encoding="utf-8")
    return output_path
